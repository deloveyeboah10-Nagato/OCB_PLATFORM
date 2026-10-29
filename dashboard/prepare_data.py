"""Create privacy-minimised dashboard CSV snapshots from approved monitoring exports.

Place these files in source_exports/ before running:
  ananse_MONITORING.xlsx
  sikacredit_MONITORING.xlsx
  oman_monitoring.xlsx

Then run: .venv\\Scripts\\python.exe prepare_data.py
"""
from pathlib import Path
import json
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "source_exports"
OUT = ROOT / "data"
SECTIONS = {
    "ananse_MONITORING.xlsx": {
        "ananse_monthly": 2, "ananse_region": 42, "ananse_channel": 1733,
        "ananse_outcomes": 2064, "ananse_customers": 2103, "ananse_top10": 5102,
    },
    "sikacredit_MONITORING.xlsx": {
        "sika_post_maturity_borrowers": 2, "sika_outstanding_borrowers": 455,
        "sika_principal_borrowers": 944, "sika_flows": 1433, "sika_monthly": 1470,
        "sika_top_post_maturity": 1508, "sika_top_principal": 1522, "sika_top_outstanding": 1537,
    },
    "oman_monitoring.xlsx": {
        "oman_customers": 2, "oman_monthly": 502, "oman_channel": 542,
        "oman_origin_country": 726, "oman_region": 1088, "oman_type": 2598,
        "oman_outcomes": 2675, "oman_top10": 2715,
    },
}
PII_COLUMNS = {"first_name", "last_name", "date_of_birth", "phone_number", "email"}


def read_section(rows, header_row):
    header = [str(v).strip() if v is not None else None for v in rows[header_row - 1]]
    while header and header[-1] is None:
        header.pop()
    records = []
    for row in rows[header_row:]:
        values = list(row[:len(header)])
        if not any(v is not None for v in values):
            break
        non_null = [v for v in values if v is not None]
        if len(non_null) == 1 and isinstance(non_null[0], str) and records:
            break
        values += [None] * (len(header) - len(values))
        records.append(values[:len(header)])
    frame = pd.DataFrame(records, columns=header)
    frame = frame.drop(columns=[c for c in frame.columns if c in PII_COLUMNS], errors="ignore")
    if "year_month" in frame.columns:
        frame["year_month"] = frame["year_month"].astype(str).str.slice(0, 7)
    return frame


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    manifest = {
        "source": "OCB monitoring workbook exports",
        "privacy_note": "Names, dates of birth, phone numbers, and email addresses are excluded from dashboard datasets.",
        "tables": {},
    }
    missing = [name for name in SECTIONS if not (SOURCE / name).exists()]
    if missing:
        raise FileNotFoundError("Missing workbook(s) in source_exports/: " + ", ".join(missing))
    for workbook_name, sections in SECTIONS.items():
        ws = load_workbook(SOURCE / workbook_name, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        for table_name, header_row in sections.items():
            frame = read_section(rows, header_row)
            output = OUT / f"{table_name}.csv"
            frame.to_csv(output, index=False)
            manifest["tables"][table_name] = {
                "source_workbook": workbook_name,
                "source_header_row": header_row,
                "rows": len(frame),
                "columns": list(frame.columns),
                "file": f"data/{table_name}.csv",
            }
    (ROOT / "data_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(f"Wrote {len(manifest['tables'])} dashboard datasets to {OUT}")


if __name__ == "__main__":
    main()
